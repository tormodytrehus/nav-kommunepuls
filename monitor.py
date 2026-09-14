#!/usr/bin/env python3
"""Publish one RSS item when NAV releases a new municipal labour-market workbook."""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import tempfile
from datetime import datetime, timezone
from email.utils import format_datetime
from html import escape
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from urllib.parse import urljoin
from xml.etree import ElementTree as ET

import requests
from openpyxl import load_workbook


USER_AGENT = "Avisa-Sor-Trondelag-NAV-monitor/1.0"
MONTHS = {
    "januar": 1,
    "februar": 2,
    "mars": 3,
    "april": 4,
    "mai": 5,
    "juni": 6,
    "juli": 7,
    "august": 8,
    "september": 9,
    "oktober": 10,
    "november": 11,
    "desember": 12,
}


def compact(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


class LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[dict[str, str]] = []
        self._href: str | None = None
        self._text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.casefold() != "a":
            return
        self._href = dict(attrs).get("href")
        self._text = []

    def handle_data(self, data: str) -> None:
        if self._href is not None:
            self._text.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.casefold() == "a" and self._href is not None:
            self.links.append({"href": self._href, "text": compact("".join(self._text))})
            self._href = None
            self._text = []


def period_from_link(text: str, href: str) -> tuple[str, str]:
    text_match = re.search(
        r"(januar|februar|mars|april|mai|juni|juli|august|september|oktober|november|desember)\s+(20\d{2})",
        text,
        re.I,
    )
    if text_match:
        month_name = text_match.group(1).casefold()
        year = int(text_match.group(2))
        month = MONTHS[month_name]
        return f"{year:04d}-{month:02d}", f"{month_name.capitalize()} {year}"
    file_match = re.search(r"(?:^|/)(20\d{2})(0[1-9]|1[0-2])[^/]*\.xlsx", href, re.I)
    if not file_match:
        raise ValueError(f"Fant ikke måned i NAV-lenken: {href}")
    year, month = int(file_match.group(1)), int(file_match.group(2))
    month_name = next(name for name, number in MONTHS.items() if number == month)
    return f"{year:04d}-{month:02d}", f"{month_name.capitalize()} {year}"


def find_latest_workbook(html: str, source_page: str) -> dict[str, str]:
    parser = LinkParser()
    parser.feed(html)
    candidates: list[dict[str, str]] = []
    for link in parser.links:
        label = link["text"].casefold()
        href = link["href"]
        if not href.casefold().split("?")[0].endswith(".xlsx"):
            continue
        if "arbeidssøkere og ledige stillinger" not in label:
            continue
        if "kommune og kjennetegn" not in label:
            continue
        period, period_label = period_from_link(link["text"], href)
        candidates.append(
            {
                "period": period,
                "period_label": period_label,
                "url": urljoin(source_page, href),
                "label": link["text"],
            }
        )
    if not candidates:
        raise RuntimeError("Fant ingen kommunefil fra NAV på kildesiden")
    return max(candidates, key=lambda item: item["period"])


def normalize_number(value: Any) -> int | float | str | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    value = compact(value)
    if not value:
        return None
    if value == "*":
        return "*"
    try:
        number = float(value.replace(",", "."))
        return int(number) if number.is_integer() else number
    except ValueError:
        return value


def extract_values(workbook_bytes: bytes, municipalities: dict[str, str]) -> dict[str, Any]:
    workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=True, read_only=True)
    if "Oversikt" not in workbook.sheetnames:
        raise RuntimeError("NAV-filen mangler arket 'Oversikt'")
    sheet = workbook["Oversikt"]
    values: dict[str, Any] = {}
    for row in sheet.iter_rows(values_only=True):
        municipality_cell = compact(row[1] if len(row) > 1 else "")
        match = re.match(r"^(\d{4})\s+(.+)$", municipality_cell)
        if not match or match.group(1) not in municipalities:
            continue
        code = match.group(1)
        values[code] = {
            "name": municipalities[code],
            "total_jobseekers": normalize_number(row[2]),
            "fully_unemployed": normalize_number(row[3]),
            "partly_unemployed": normalize_number(row[4]),
            "labour_market_measures": normalize_number(row[5]),
            "fully_unemployed_pct": normalize_number(row[6]),
            "partly_unemployed_pct": normalize_number(row[7]),
            "measures_pct": normalize_number(row[8]),
            "new_vacancies": normalize_number(row[9]),
        }
    missing = [f"{code} {name}" for code, name in municipalities.items() if code not in values]
    if missing:
        raise RuntimeError(f"Fant ikke alle kommunene i NAV-filen: {', '.join(missing)}")
    return values


def display(value: Any, percent: bool = False) -> str:
    if value is None:
        return "ikke oppgitt"
    if value == "*":
        return "skjermet av NAV"
    rendered = str(value).replace(".", ",")
    return f"{rendered} %" if percent else rendered


def change_text(current: Any, previous: Any) -> str:
    if not isinstance(current, (int, float)) or not isinstance(previous, (int, float)):
        return ""
    difference = current - previous
    if difference == 0:
        return " (uendret)"
    sign = "+" if difference > 0 else ""
    return f" ({sign}{difference:g} fra forrige måned)".replace(".", ",")


def make_event(
    workbook: dict[str, str],
    values: dict[str, Any],
    previous_values: dict[str, Any],
    observed_at: datetime,
) -> dict[str, Any]:
    return {
        "id": f"nav-kommune-{workbook['period']}",
        "period": workbook["period"],
        "period_label": workbook["period_label"],
        "source_url": workbook["url"],
        "observed_at": observed_at.astimezone(timezone.utc).isoformat(),
        "values": values,
        "previous_values": previous_values,
    }


def build_rss(events: list[dict[str, Any]], config: dict[str, Any]) -> bytes:
    source_page = config["source_page"]
    rss = ET.Element("rss", {"version": "2.0"})
    channel = ET.SubElement(rss, "channel")
    ET.SubElement(channel, "title").text = config["feed"]["title"]
    ET.SubElement(channel, "link").text = source_page
    ET.SubElement(channel, "description").text = config["feed"]["description"]
    ET.SubElement(channel, "language").text = "nb-NO"

    for event in events:
        item = ET.SubElement(channel, "item")
        ET.SubElement(item, "title").text = f"NAV-kommunepuls: {event['period_label']}"
        ET.SubElement(item, "link").text = source_page
        guid = ET.SubElement(item, "guid", {"isPermaLink": "false"})
        guid.text = event["id"]
        ET.SubElement(item, "pubDate").text = format_datetime(
            datetime.fromisoformat(event["observed_at"])
        )
        rows: list[str] = []
        for code, current in event["values"].items():
            previous = event.get("previous_values", {}).get(code, {})
            unemployed_change = change_text(
                current.get("fully_unemployed"), previous.get("fully_unemployed")
            )
            rows.append(
                f"<strong>{escape(current['name'])}</strong>: "
                f"{display(current.get('fully_unemployed'))} helt ledige"
                f"{escape(unemployed_change)}; "
                f"andel {display(current.get('fully_unemployed_pct'), percent=True)}; "
                f"{display(current.get('new_vacancies'))} nye stillinger"
            )
        rows.append(
            f'<a href="{escape(event["source_url"])}">Originaltabell fra NAV (Excel)</a>'
        )
        rows.append(f'<a href="{escape(source_page)}">NAVs statistikkside</a>')
        ET.SubElement(item, "description").text = "<br>".join(rows)

    ET.indent(rss)
    return ET.tostring(rss, encoding="utf-8", xml_declaration=True)


def load_json(path: Path, default: dict[str, Any]) -> dict[str, Any]:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def atomic_write(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as temporary:
            temporary.write(data)
        os.replace(temporary_name, path)
    finally:
        if os.path.exists(temporary_name):
            os.unlink(temporary_name)


def run(config_path: Path, state_path: Path, feed_path: Path) -> int:
    config = load_json(config_path, {})
    previous = load_json(
        state_path,
        {
            "version": 1,
            "initialized": False,
            "latest_period": None,
            "latest_source_url": None,
            "latest_values": {},
            "events": [],
        },
    )
    session = requests.Session()
    session.headers["User-Agent"] = USER_AGENT
    source_response = session.get(config["source_page"], timeout=60)
    source_response.raise_for_status()
    workbook = find_latest_workbook(source_response.text, config["source_page"])

    if previous.get("latest_period") == workbook["period"]:
        atomic_write(feed_path, build_rss(previous.get("events", []), config))
        print(f"Ingen ny NAV-publisering etter {workbook['period_label']}")
        return 0

    workbook_response = session.get(workbook["url"], timeout=120)
    workbook_response.raise_for_status()
    values = extract_values(workbook_response.content, config["municipalities"])
    observed_at = datetime.now(timezone.utc)
    events = list(previous.get("events", []))
    if previous.get("initialized") or not config.get("silent_first_run", True):
        events.insert(
            0,
            make_event(
                workbook,
                values,
                previous.get("latest_values", {}),
                observed_at,
            ),
        )
    events = events[: int(config.get("max_feed_items", 24))]
    state = {
        "version": 1,
        "initialized": True,
        "latest_period": workbook["period"],
        "latest_source_url": workbook["url"],
        "latest_values": values,
        "events": events,
    }
    atomic_write(state_path, json.dumps(state, ensure_ascii=False, indent=2).encode() + b"\n")
    atomic_write(feed_path, build_rss(events, config))
    print(f"Behandlet NAV-publisering for {workbook['period_label']}; {len(events)} RSS-innlegg")
    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default="config.json")
    parser.add_argument("--state", default="state/nav.json")
    parser.add_argument("--feed", default="public/feed.xml")
    return parser.parse_args()


if __name__ == "__main__":
    arguments = parse_args()
    raise SystemExit(run(Path(arguments.config), Path(arguments.state), Path(arguments.feed)))
